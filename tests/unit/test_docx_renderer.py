from __future__ import annotations

import base64
from pathlib import Path

from docx import Document

from mcp_market_research.models import (
    Question,
    QuestionType,
    RenderRequest,
    Section,
)
from mcp_market_research.renderers.docx_renderer import render


def _request() -> RenderRequest:
    return RenderRequest(
        title="Customer Experience Survey",
        language="en",
        country_code="US",
        company_name="Acme Corp",
        sections=[
            Section(
                heading="Awareness",
                questions=[
                    Question(text="How did you hear about us?", type=QuestionType.OPEN_TEXT, required=True),
                    Question(
                        text="Which platforms do you use?",
                        type=QuestionType.MULTI_CHOICE,
                        options=["Instagram", "TikTok", "YouTube"],
                        required=True,
                    ),
                ],
            ),
            Section(
                heading="Satisfaction",
                questions=[
                    Question(
                        text="Rate your experience.", type=QuestionType.RATING,
                        options=["1", "2", "3", "4", "5"],
                    ),
                ],
            ),
        ],
    )


def test_renders_questionnaire(sample_docx_path: Path, output_dir: Path) -> None:
    result = render(sample_docx_path, _request(), output_dir)
    out = Path(result.path)
    assert out.exists()
    assert out.suffix == ".docx"
    assert result.size_bytes > 0
    assert result.bytes_b64 is not None

    decoded = base64.b64decode(result.bytes_b64)
    assert decoded[:2] == b"PK"  # zip magic — .docx is a zip


def test_substitutes_placeholders(sample_docx_path: Path, output_dir: Path) -> None:
    result = render(sample_docx_path, _request(), output_dir)
    doc = Document(result.path)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Acme Corp" in text
    assert "Customer Experience Survey" in text
    assert "Awareness" in text
    assert "How did you hear about us?" in text
    assert "TikTok" in text
    assert "{{" not in text  # all placeholders substituted


def test_filename_includes_language_and_slug(sample_docx_path: Path, output_dir: Path) -> None:
    result = render(sample_docx_path, _request(), output_dir)
    assert result.filename.startswith("questionnaire_acme-corp_en_")
    assert result.filename.endswith(".docx")
