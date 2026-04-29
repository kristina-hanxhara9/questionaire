"""End-to-end: parse Excel + PDF, compose payload, render .docx, verify content."""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest
from docx import Document

from mcp_market_research.models import (
    Question,
    QuestionType,
    RenderRequest,
    Section,
)
from mcp_market_research.tools.channels import (
    get_channel_guide_tool,
    list_channels_tool,
)
from mcp_market_research.tools.render import render_questionnaire_docx_tool
from mcp_market_research.tools.template import get_template_structure_tool


def _payload_from_guide(guide: Any, language: str, country: str | None) -> RenderRequest:
    sections: list[Section] = []
    for section in guide.sections:
        questions = []
        for row in section.rows:
            options: list[str] | None = None
            if row.question_type in {QuestionType.SINGLE_CHOICE, QuestionType.MULTI_CHOICE}:
                options = ["Option A", "Option B", "Option C"]
            elif row.question_type == QuestionType.RATING:
                options = ["1", "2", "3", "4", "5"]
            elif row.question_type == QuestionType.YES_NO:
                options = ["Yes", "No"]
            text = row.example or row.guidance
            questions.append(
                Question(
                    text=text,
                    type=row.question_type,
                    options=options,
                    required=row.required,
                )
            )
        sections.append(Section(heading=section.heading, questions=questions))
    return RenderRequest(
        title="Customer Experience Survey",
        language=language,
        country_code=country,
        company_name="Acme Corporation",
        sections=sections,
    )


@pytest.mark.integration
def test_full_flow_english(settings: Any) -> None:
    channels = list_channels_tool(settings)
    assert "social_media" in channels

    guide = get_channel_guide_tool(settings, "social_media")
    structure = get_template_structure_tool(settings)
    assert structure.placeholders

    payload = _payload_from_guide(guide, language="en", country="US")
    result = render_questionnaire_docx_tool(settings, payload)
    out = Path(result.path)
    assert out.exists()
    text = "\n".join(p.text for p in Document(out).paragraphs)
    assert "Acme Corporation" in text
    assert "Customer Experience Survey" in text


@pytest.mark.integration
def test_full_flow_french_in_store(settings: Any) -> None:
    guide = get_channel_guide_tool(settings, "in_store")
    payload = _payload_from_guide(guide, language="fr", country="FR")
    result = render_questionnaire_docx_tool(settings, payload)
    assert Path(result.path).exists()
    assert result.filename.startswith("questionnaire_acme-corporation_fr_")


@pytest.mark.integration
def test_full_flow_german_email(settings: Any) -> None:
    guide = get_channel_guide_tool(settings, "email")
    payload = _payload_from_guide(guide, language="de", country="DE")
    result = render_questionnaire_docx_tool(settings, payload)
    assert Path(result.path).exists()
    assert result.filename.startswith("questionnaire_acme-corporation_de_")
